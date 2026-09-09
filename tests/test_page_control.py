import io
import os
import zipfile
import unittest
from pathlib import Path
from PIL import Image
import pypdf
import docx

from config import settings
from services.document import DocumentService, DocumentSecurityError


class TestPageControlAndMultiFormat(unittest.TestCase):
    def setUp(self):
        self.spool_dir = Path("data/spool_test_pages")
        self.spool_dir.mkdir(parents=True, exist_ok=True)
        self.orig_spool = settings.SPOOL_DIR
        settings.SPOOL_DIR = self.spool_dir

    def tearDown(self):
        settings.SPOOL_DIR = self.orig_spool
        for f in self.spool_dir.glob("*"):
            try:
                f.unlink()
            except Exception:
                pass
        try:
            self.spool_dir.rmdir()
        except Exception:
            pass

    def _create_test_pdf(self, num_pages: int, width: float = 595.28, height: float = 841.89) -> bytes:
        writer = pypdf.PdfWriter()
        for i in range(num_pages):
            writer.add_blank_page(width=width, height=height)
        buf = io.BytesIO()
        writer.write(buf)
        return buf.getvalue()

    def test_parse_page_range_strictly(self):
        """Проверка строгого парсинга диапазонов и исключения дубликатов"""
        norm, count = DocumentService.parse_page_range("1-3, 5, 7-8", total_pages=10)
        self.assertEqual(norm, "1-3, 5, 7-8")
        self.assertEqual(count, 6)

        # Дубликаты должны схлопываться
        norm, count = DocumentService.parse_page_range("1, 1, 1-2, 2", total_pages=5)
        self.assertEqual(norm, "1-2")
        self.assertEqual(count, 2)

        # "Все" страницы
        norm, count = DocumentService.parse_page_range("all", total_pages=5)
        self.assertEqual(norm, "all")
        self.assertEqual(count, 5)

        # Выход за пределы документа
        with self.assertRaises(DocumentSecurityError):
            DocumentService.parse_page_range("1-6", total_pages=5)

        with self.assertRaises(DocumentSecurityError):
            DocumentService.parse_page_range("0-3", total_pages=5)

        # Обратный диапазон
        with self.assertRaises(DocumentSecurityError):
            DocumentService.parse_page_range("5-2", total_pages=10)

    def test_physical_pdf_slicing(self):
        """Физическая нарезка страниц: файл на диске должен содержать ТОЛЬКО запрошенные страницы"""
        pdf_bytes = self._create_test_pdf(6)
        src_path = self.spool_dir / "test_orig.pdf"
        src_path.write_bytes(pdf_bytes)

        # Запрашиваем страницы 2, 4-5 (всего 3 страницы)
        prepared_path = DocumentService.prepare_job_pdf(
            source_pdf=src_path,
            order_uuid="slice_test_1",
            selected_pages="2, 4-5",
            copies=1
        )

        self.assertTrue(prepared_path.exists())
        reader = pypdf.PdfReader(str(prepared_path))
        self.assertEqual(len(reader.pages), 3, "Нарезанный PDF должен содержать ровно 3 страницы!")

    def test_copies_duplication_for_raw_socket(self):
        """Для сокетной печати (RAW 9100) страницы тиражируются в PDF на нужное число копий"""
        pdf_bytes = self._create_test_pdf(4)
        src_path = self.spool_dir / "test_copies.pdf"
        src_path.write_bytes(pdf_bytes)

        # Запрашиваем страницы 1-2 с 3 копиями -> должно быть 6 страниц
        prepared_path = DocumentService.prepare_job_pdf(
            source_pdf=src_path,
            order_uuid="copies_test_1",
            selected_pages="1-2",
            copies=3
        )

        reader = pypdf.PdfReader(str(prepared_path))
        self.assertEqual(len(reader.pages), 6, "2 страницы x 3 копии = 6 страниц в потоке печати")

    def test_a4_page_geometry_normalization(self):
        """Масштабирование нестандартных страниц (A3 / баннеры) под формат A4"""
        # Создаем PDF размером 1500 x 2500 pt
        huge_pdf = self._create_test_pdf(1, width=1500, height=2500)
        src_path = self.spool_dir / "huge.pdf"
        src_path.write_bytes(huge_pdf)

        prepared_path = DocumentService.prepare_job_pdf(
            source_pdf=src_path,
            order_uuid="huge_test",
            selected_pages="all",
            copies=1
        )

        reader = pypdf.PdfReader(str(prepared_path))
        page = reader.pages[0]
        # Проверяем, что mediabox смасштабирован до формата A4 (595.28 x 841.89)
        self.assertAlmostEqual(float(page.mediabox.width), 595.28, delta=1.0)
        self.assertAlmostEqual(float(page.mediabox.height), 841.89, delta=1.0)

    def test_txt_to_a4_pdf_pagination(self):
        """Текстовые файлы (.txt) разбиваются на страницы с нумерацией и колонтитулами"""
        # 120 строк при 48 строках на страницу -> ровно 3 страницы A4
        lines = [f"Лабораторная работа #1, строка {i} для студента ЦСО-4" for i in range(120)]
        txt_bytes = "\n".join(lines).encode("utf-8")

        dest_pdf = self.spool_dir / "text_doc.pdf"
        pages = DocumentService.convert_txt_to_a4_pdf(txt_bytes, dest_pdf)

        self.assertEqual(pages, 3, "120 строк текста должны сформировать ровно 3 страницы A4")
        self.assertTrue(dest_pdf.exists())

        reader = pypdf.PdfReader(str(dest_pdf))
        self.assertEqual(len(reader.pages), 3)

    def test_multi_frame_image_to_multipage_pdf(self):
        """Многокадровые изображения (TIFF) конвертируются в многостраничный PDF"""
        # Создаем 3-кадровый TIFF
        frames = [
            Image.new("RGB", (800, 600), (255, 0, 0)),
            Image.new("RGB", (800, 600), (0, 255, 0)),
            Image.new("RGB", (800, 600), (0, 0, 255))
        ]
        tiff_buf = io.BytesIO()
        frames[0].save(tiff_buf, format="TIFF", save_all=True, append_images=frames[1:])
        tiff_bytes = tiff_buf.getvalue()

        dest_pdf = self.spool_dir / "multiframe.pdf"
        pages = DocumentService.convert_image_to_a4_pdf(tiff_bytes, dest_pdf)

        self.assertEqual(pages, 3, "3-кадровый TIFF должен сформировать 3 страницы PDF")
        reader = pypdf.PdfReader(str(dest_pdf))
        self.assertEqual(len(reader.pages), 3)

    def test_docx_zip_bomb_protection(self):
        """Защита от Zip-бомбы в файле .docx"""
        import services.document as doc_mod
        orig_max = doc_mod.MAX_ZIP_UNCOMPRESSED_BYTES
        doc_mod.MAX_ZIP_UNCOMPRESSED_BYTES = 500  # Снижаем лимит для теста
        try:
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w") as z:
                z.writestr("word/document.xml", b"<xml>" + b"A" * 1000 + b"</xml>")

            with self.assertRaises(DocumentSecurityError) as ctx:
                DocumentService.convert_docx_to_pdf(buf.getvalue(), self.spool_dir / "bomb.pdf", "bomb_uuid")
            self.assertIn("Zip-бомб", str(ctx.exception))
        finally:
            doc_mod.MAX_ZIP_UNCOMPRESSED_BYTES = orig_max

    def test_docx_text_fallback_conversion(self):
        """Автономная конвертация .docx в A4 PDF при отсутствии LibreOffice"""
        doc = docx.Document()
        doc.add_heading("Отчет по практике в ЦСО-4", level=1)
        doc.add_paragraph("Принтер Pantum BP2300NW настроен для работы в сети роутера Huawei AX3.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Параметр"
        table.cell(0, 1).text = "Значение"
        table.cell(1, 0).text = "Статус"
        table.cell(1, 1).text = "Online"

        doc_buf = io.BytesIO()
        doc.save(doc_buf)
        docx_bytes = doc_buf.getvalue()

        dest_pdf = self.spool_dir / "docx_test.pdf"
        pages = DocumentService.convert_docx_to_pdf(docx_bytes, dest_pdf, "docx_uuid_1")

        self.assertGreaterEqual(pages, 1)
        self.assertTrue(dest_pdf.exists())
        reader = pypdf.PdfReader(str(dest_pdf))
        self.assertEqual(len(reader.pages), pages)

    def test_corrupt_pdf_pages_detection(self):
        """Обнаружение PDF с поврежденными или нулевыми размерами страниц"""
        # Создаем PDF со страницей нулевого размера
        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=0, height=0)
        buf = io.BytesIO()
        writer.write(buf)

        corrupt_path = self.spool_dir / "corrupt_dims.pdf"
        corrupt_path.write_bytes(buf.getvalue())

        with self.assertRaises(DocumentSecurityError):
            DocumentService.verify_pdf(corrupt_path)


if __name__ == "__main__":
    unittest.main()
